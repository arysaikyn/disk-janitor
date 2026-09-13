#!/usr/bin/env python3
"""Generate the 开销明细表 — AI development cost by phase (作业要求 item 4).

Writes BOTH artifacts from one mined snapshot so they can never disagree:
  * docs/expense-sheet.xlsx          — the graded Excel deliverable
  * docs/expense-sheet.md            — the tables between the GENERATED markers

Source of truth: the omp session transcripts for this repo. Every assistant API call
in those files carries a `message.usage` record; tool invocations appear as `toolCall`
content items. Nothing here is hand-typed.

Usage:
    python3 scripts/gen-expense-xlsx.py                # regenerate md block + xlsx
    python3 scripts/gen-expense-xlsx.py --json         # print mined figures only
    python3 scripts/gen-expense-xlsx.py --xlsx-only    # skip the markdown update

Requires: openpyxl (only for the xlsx path).

Method (mirrored in the 方法 sheet and in docs/expense-sheet.md §3):
  * Phases key to the milestone dates recorded in docs/ai-logs/M1.md…M5.md.
  * Days are bucketed by the message timestamp's UTC date; day rows sum exactly to
    the phase rows and the grand total.
  * 人时 is DERIVED, not observed: timestamps are sorted globally, a gap > 20 minutes
    splits a segment, and a segment is attributed to the phase of its start day.
  * The snapshot cutoff is the newest timestamp in the corpus, so the output is a
    function of the corpus alone — re-running after new work moves the numbers, and
    both artifacts move together.
  * `cost.cacheWrite` is 0 on every record (this gateway bills reads only) — not a bug.
  * P3 reports $0.00 because the DeepSeek-V4.1-Flash gateway publishes no per-token
    price. Token counts are real; $0.00 is a reporting gap, not free usage.
"""

from __future__ import annotations

import argparse
import collections
import datetime
import glob
import json
import os
import sys

SESSIONS = "~/.omp/agent/sessions/-Code-disk_janitor"
USD_CNY = 7.2
GAP_SECONDS = 20 * 60
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Phase → the UTC days belonging to it, named after the milestone logs. 09-06…09-08 is
# design lock-in that precedes M1; 09-02…09-05 and 09-11 recorded no sessions.
PHASES: list[tuple[str, str, list[str]]] = [
    ("P0", "立项 · 调研与设计", ["2026-08-30", "2026-08-31", "2026-09-01",
                              "2026-09-06", "2026-09-07", "2026-09-08"]),
    ("P1", "M1+M2 核心引擎", ["2026-09-09"]),
    ("P2", "M3+M4 GUI 与 Agent", ["2026-09-10"]),
    ("P3", "M5 交付与收尾", ["2026-09-12", "2026-09-13"]),
]
DAY_TO_PHASE = {d: name for name, _, ds in PHASES for d in ds}
PHASE_LABEL = {name: f"{name} {label}" for name, label, _ in PHASES}

BEGIN = "<!-- BEGIN GENERATED:expense — edit scripts/gen-expense-xlsx.py, not this block -->"
END = "<!-- END GENERATED:expense -->"


def parse_ts(raw: str) -> datetime.datetime | None:
    try:
        return datetime.datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except (ValueError, AttributeError):
        return None


def mine(root: str) -> dict:
    root = os.path.expanduser(root)
    if not os.path.isdir(root):
        sys.exit(f"session directory not found: {root}")

    records: dict[str, dict] = {}   # dedup by message id is mandatory (files overlap)
    stamps: list[datetime.datetime] = []

    paths = sorted(glob.glob(os.path.join(root, "*.jsonl")))
    paths += sorted(glob.glob(os.path.join(root, "*/*.jsonl")))

    for path in paths:
        with open(path, errors="replace") as fh:
            for line in fh:
                # Activity stamps: every timestamp-bearing line counts. Do NOT early-exit
                # on '"usage"' here — that silently drops records for the 人时 pass.
                idx = line.find('"timestamp":"')
                if idx >= 0:
                    ts = parse_ts(line[idx + 13: idx + 37])
                    if ts is not None:
                        stamps.append(ts)

                if '"usage"' not in line:
                    continue
                try:
                    doc = json.loads(line)
                except json.JSONDecodeError:
                    continue
                msg = doc.get("message") or {}
                usage = msg.get("usage")
                if not isinstance(usage, dict):
                    continue
                ts = parse_ts(doc.get("timestamp", ""))
                if ts is None:
                    continue
                day = ts.date().isoformat()
                if day not in DAY_TO_PHASE:
                    continue

                key = doc.get("id") or f"{day}:{len(records)}"
                if key in records:
                    continue    # the same message appears in more than one transcript
                cost = usage.get("cost") or {}
                records[key] = {
                    "day": day,
                    "phase": DAY_TO_PHASE[day],
                    "model": msg.get("model") or "?",
                    "tools": [i.get("name") or "?"
                              for i in msg.get("content") or []
                              if isinstance(i, dict) and i.get("type") == "toolCall"],
                    "input": usage.get("input") or 0,
                    "output": usage.get("output") or 0,
                    "cacheRead": usage.get("cacheRead") or 0,
                    "cacheWrite": usage.get("cacheWrite") or 0,
                    "usd": cost.get("total") or 0.0,
                }

    # Model and tool histograms are built from the DEDUPLICATED records: counting them
    # where they are parsed would multiply every message by the number of transcripts
    # that mention it.
    models: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    tool_calls: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    for rec in records.values():
        models[rec["phase"]][rec["model"]] += 1
        for tool in rec["tools"]:
            tool_calls[rec["phase"]][tool] += 1

    # 人时: global union of activity windows; >20 min gap splits a segment.
    stamps = sorted(set(stamps))
    segments: list[tuple[datetime.datetime, datetime.datetime]] = []
    if stamps:
        start = prev = stamps[0]
        for ts in stamps[1:]:
            if (ts - prev).total_seconds() > GAP_SECONDS:
                segments.append((start, prev))
                start = ts
            prev = ts
        segments.append((start, prev))

    return {"records": records, "models": models, "tools": tool_calls,
            "segments": segments,
            "sessions": len(glob.glob(os.path.join(root, "*.jsonl"))),
            "subagents": len(glob.glob(os.path.join(root, "*/*.jsonl"))),
            "cutoff": max(stamps) if stamps else None,
            "first": min(stamps) if stamps else None}


def rollup(mined: dict) -> dict:
    per_day: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    per_phase: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    for rec in mined["records"].values():
        for bucket in (per_day[rec["day"]], per_phase[rec["phase"]]):
            bucket["calls"] += 1
            bucket["input"] += rec["input"]
            bucket["output"] += rec["output"]
            bucket["cacheRead"] += rec["cacheRead"]
            bucket["usd"] += rec["usd"]

    hours_day: dict[str, float] = collections.defaultdict(float)
    segs_day: collections.Counter = collections.Counter()
    for a, b in mined["segments"]:
        day = a.date().isoformat()
        hours_day[day] += (b - a).total_seconds() / 3600
        segs_day[day] += 1

    hours_phase: dict[str, float] = collections.defaultdict(float)
    segs_phase: collections.Counter = collections.Counter()
    for day, hours in hours_day.items():
        phase = DAY_TO_PHASE.get(day)
        if phase:
            hours_phase[phase] += hours
            segs_phase[phase] += segs_day[day]

    return {"per_day": per_day, "per_phase": per_phase, "hours_day": hours_day,
            "segs_day": segs_day, "hours_phase": hours_phase, "segs_phase": segs_phase}


def totals_of(rolled: dict) -> collections.Counter:
    totals: collections.Counter = collections.Counter()
    for bucket in rolled["per_phase"].values():
        totals.update(bucket)
    return totals


def fmt_int(n) -> str:
    return f"{int(n):,}"


# --------------------------------------------------------------------------- markdown

def render_markdown(mined: dict, rolled: dict) -> str:
    grand_hours = sum(rolled["hours_phase"].values())
    total = totals_of(rolled)
    total_tokens = total["input"] + total["output"] + total["cacheRead"]
    cutoff = "?"
    if mined["cutoff"]:
        stamp_ts = mined["cutoff"]
        local = stamp_ts.astimezone(datetime.timezone(datetime.timedelta(hours=8)))
        cutoff = (f"{local:%Y-%m-%d %H:%M} +08:00 "
                  f"(= {stamp_ts.astimezone(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC)")

    L: list[str] = []
    L.append(f"*Snapshot cutoff: **{cutoff}** — the newest activity in the corpus. "
             f"Corpus: {mined['sessions']} top-level sessions + {mined['subagents']} "
             f"subagent transcripts. Both this block and the `.xlsx` are emitted by the "
             f"same run of `scripts/gen-expense-xlsx.py`.*")
    L.append("")

    L.append("| 阶段 phase | 起止 dates | **人时 (h)** | **API次数 calls** | **输入 tokens** | "
             "**输出 tokens** | **cache read** | **总 tokens** | **花费 USD** | **花费 CNY** |")
    L.append("|---|---|---|---|---|---|---|---|---|---|")
    for name, _, days in PHASES:
        b = rolled["per_phase"].get(name, collections.Counter())
        h = rolled["hours_phase"].get(name, 0.0)
        tt = b["input"] + b["output"] + b["cacheRead"]
        L.append(f"| **{PHASE_LABEL[name]}** | {days[0][5:]} … {days[-1][5:]} | {h:.2f} | "
                 f"{fmt_int(b['calls'])} | {fmt_int(b['input'])} | {fmt_int(b['output'])} | "
                 f"{fmt_int(b['cacheRead'])} | {fmt_int(tt)} | ${b['usd']:.4f} | "
                 f"¥{b['usd'] * USD_CNY:.2f} |")
    L.append(f"| **合计 TOTAL** | **{PHASES[0][2][0][5:]} … {PHASES[-1][2][-1][5:]}** | "
             f"**{grand_hours:.2f}** | **{fmt_int(total['calls'])}** | "
             f"**{fmt_int(total['input'])}** | **{fmt_int(total['output'])}** | "
             f"**{fmt_int(total['cacheRead'])}** | **{fmt_int(total_tokens)}** | "
             f"**${total['usd']:.4f}** | **¥{total['usd'] * USD_CNY:.2f}** |")
    L.append("")

    L.append("### 模型 models (API calls per phase)")
    L.append("")
    L.append("| 阶段 | 模型 (API calls) |")
    L.append("|---|---|")
    for name, _, _ in PHASES:
        row = " · ".join(f"`{m}` {fmt_int(c)}"
                         for m, c in mined["models"].get(name, {}).most_common())
        L.append(f"| {name} | {row} |")
    L.append("")

    L.append("### 工具 tools (tool invocations per phase)")
    L.append("")
    L.append("| 阶段 | 调用总数 | top tools |")
    L.append("|---|---|---|")
    all_tools: collections.Counter = collections.Counter()
    for counter in mined["tools"].values():
        all_tools.update(counter)
    for name, _, _ in PHASES:
        counter = mined["tools"].get(name, collections.Counter())
        top = " · ".join(f"`{t}` {fmt_int(c)}" for t, c in counter.most_common(6))
        L.append(f"| {name} | {fmt_int(sum(counter.values()))} | {top} |")
    top_all = " · ".join(f"`{t}` {fmt_int(c)}" for t, c in all_tools.most_common(6))
    L.append(f"| **合计** | **{fmt_int(sum(all_tools.values()))}** | {top_all} |")
    L.append("")

    L.append("### 逐日 daily")
    L.append("")
    L.append("| 日期 | 阶段 | 人时 h | calls | input | output | cache read | USD |")
    L.append("|---|---|---|---|---|---|---|---|")
    for day in sorted(rolled["per_day"]):
        b = rolled["per_day"][day]
        h = rolled["hours_day"].get(day, 0.0)
        L.append(f"| {day[5:]} | {DAY_TO_PHASE[day]} | {h:.2f} | {fmt_int(b['calls'])} | "
                 f"{fmt_int(b['input'])} | {fmt_int(b['output'])} | {fmt_int(b['cacheRead'])} | "
                 f"{b['usd']:.4f} |")
    L.append("")

    L.append("### 人时 person-hours")
    L.append("")
    L.append(f"Union of all activity windows across the transcripts: timestamps sorted, a gap "
             f"**> {GAP_SECONDS // 60} minutes** splits a segment, segments merged globally so "
             f"overlapping sessions count once, and each segment attributed to the phase of its "
             f"start day.")
    L.append("")
    L.append("| 阶段 | 人时 h | segments |")
    L.append("|---|---|---|")
    for name, _, _ in PHASES:
        L.append(f"| {name} | {rolled['hours_phase'].get(name, 0.0):.2f} | "
                 f"{rolled['segs_phase'].get(name, 0)} |")
    L.append(f"| **合计** | **{grand_hours:.2f}** | **{len(mined['segments'])}** |")
    L.append("")
    span = ""
    if mined["first"] and mined["cutoff"]:
        wall = (mined["cutoff"] - mined["first"]).total_seconds() / 3600
        span = (f"Raw wall-clock span {mined['first']:%m-%d %H:%M} → {mined['cutoff']:%m-%d %H:%M} "
                f"UTC ≈ {wall:.0f} h — **not** the number; it includes sleep, class, and other work.")
    L.append(f"Exact unrounded total: **{grand_hours:.4f} h**. " + span)
    L.append("")

    L.append("### 会话 session metrics")
    L.append("")
    L.append("| Metric | Value | Rule |")
    L.append("|---|---|---|")
    L.append(f"| Sessions | **{mined['sessions']}** | top-level transcripts under `~/.omp/agent/sessions/-Code-disk_janitor/` |")
    L.append(f"| Subagent transcripts | {mined['subagents']} | nested session files (parallel fan-out) |")
    L.append(f"| API calls | **{fmt_int(total['calls'])}** | billable assistant round-trips |")
    L.append(f"| 人时 | **{grand_hours:.2f} h** | union of activity windows, {len(mined['segments'])} segments |")
    L.append(f"| Elapsed | **15 days** | 08-30 → 09-13 inclusive |")
    return "\n".join(L)


def update_markdown(path: str, block: str) -> None:
    with open(path) as fh:
        text = fh.read()
    i, j = text.find(BEGIN), text.find(END)
    if i < 0 or j < 0 or j < i:
        sys.exit(f"markers not found in {path} — expected BEGIN/END GENERATED comments")
    new = text[: i + len(BEGIN)] + "\n\n" + block + "\n\n" + text[j:]
    with open(path, "w") as fh:
        fh.write(new)


# ------------------------------------------------------------------------------- xlsx

def build_workbook(mined: dict, rolled: dict, out_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl not installed: python3 -m pip install openpyxl")

    wb = Workbook()
    # Pin the document timestamps so the emitted bytes depend only on the corpus
    # (openpyxl otherwise stamps "now" into docProps/core.xml on every run).
    stamp = (mined["cutoff"] or datetime.datetime(2026, 9, 13, tzinfo=datetime.timezone.utc))
    wb.properties.created = stamp.replace(tzinfo=None)
    wb.properties.modified = stamp.replace(tzinfo=None)
    wb.properties.creator = "scripts/gen-expense-xlsx.py"
    wb.properties.lastModifiedBy = "scripts/gen-expense-xlsx.py"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    bold = Font(bold=True)
    grand_hours = sum(rolled["hours_phase"].values())
    total = totals_of(rolled)
    total_tokens = total["input"] + total["output"] + total["cacheRead"]

    def style_header(ws, width: int) -> None:
        for col in range(1, width + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def autosize(ws, cap: int = 52) -> None:
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, cap)

    ws = wb.active
    ws.title = "分阶段 by phase"
    ws.append(["阶段 phase", "起止 dates", "人时 (h)", "API次数 calls", "输入 tokens",
               "输出 tokens", "cache read", "总 tokens", "花费 USD", "花费 CNY",
               "模型 models", "工具 tools"])
    style_header(ws, 12)
    for name, _, days in PHASES:
        b = rolled["per_phase"].get(name, collections.Counter())
        tt = b["input"] + b["output"] + b["cacheRead"]
        models = " · ".join(f"{m} {c}" for m, c in mined["models"].get(name, {}).most_common())
        tools = " · ".join(f"{t} {c}" for t, c in mined["tools"].get(name, {}).most_common(6))
        ws.append([PHASE_LABEL[name], f"{days[0]} … {days[-1]}",
                   round(rolled["hours_phase"].get(name, 0.0), 2), b["calls"], b["input"],
                   b["output"], b["cacheRead"], tt, round(b["usd"], 4),
                   round(b["usd"] * USD_CNY, 4), models, tools])
    ws.append(["合计 TOTAL", f"{PHASES[0][2][0]} … {PHASES[-1][2][-1]}", round(grand_hours, 2),
               total["calls"], total["input"], total["output"], total["cacheRead"],
               total_tokens, round(total["usd"], 4), round(total["usd"] * USD_CNY, 4), "", ""])
    for cell in ws[ws.max_row]:
        cell.font = bold
    autosize(ws)

    ws2 = wb.create_sheet("逐日 daily")
    ws2.append(["日期 date", "阶段 phase", "人时 h", "API次数 calls", "输入 tokens",
                "输出 tokens", "cache read", "总 tokens", "花费 USD", "花费 CNY"])
    style_header(ws2, 10)
    for day in sorted(rolled["per_day"]):
        b = rolled["per_day"][day]
        tt = b["input"] + b["output"] + b["cacheRead"]
        ws2.append([day, DAY_TO_PHASE[day], round(rolled["hours_day"].get(day, 0.0), 2),
                    b["calls"], b["input"], b["output"], b["cacheRead"], tt,
                    round(b["usd"], 4), round(b["usd"] * USD_CNY, 4)])
    ws2.append(["合计 TOTAL", "", round(grand_hours, 2), total["calls"], total["input"],
                total["output"], total["cacheRead"], total_tokens,
                round(total["usd"], 4), round(total["usd"] * USD_CNY, 4)])
    for cell in ws2[ws2.max_row]:
        cell.font = bold
    autosize(ws2)

    ws3 = wb.create_sheet("工具 tools")
    names = [n for n, _, _ in PHASES]
    ws3.append(["工具 tool", *names, "合计 total"])
    style_header(ws3, len(names) + 2)
    all_tools: collections.Counter = collections.Counter()
    for counter in mined["tools"].values():
        all_tools.update(counter)
    for tool, _ in all_tools.most_common():
        ws3.append([tool, *[mined["tools"].get(n, {}).get(tool, 0) for n in names],
                    all_tools[tool]])
    ws3.append(["合计 total", *[sum(mined["tools"].get(n, {}).values()) for n in names],
                sum(all_tools.values())])
    for cell in ws3[ws3.max_row]:
        cell.font = bold
    autosize(ws3)

    ws4 = wb.create_sheet("方法 method")
    ws4.append(["项目 item", "说明 detail"])
    style_header(ws4, 2)
    stamp_ts = mined["cutoff"]
    local = stamp_ts.astimezone(datetime.timezone(datetime.timedelta(hours=8))) if stamp_ts else None
    cutoff = (f"{local:%Y-%m-%d %H:%M} +08:00 "
              f"({stamp_ts.astimezone(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC)") if stamp_ts else "?"
    for row in [
        ("项目 project", "Disk Janitor / 还我磁盘 Giveback"),
        ("课程 course", "程序设计训练（Rust 语言）30240522-3, THU CS"),
        ("交付物 deliverable",
         "AI 开发开销明细表（分阶段: 人时 / API次数 / tokens / 花费 / 模型 / 工具）"),
        ("数据来源 source", f"omp session transcripts: {SESSIONS}"),
        ("语料 corpus",
         f"{mined['sessions']} top-level sessions + {mined['subagents']} subagent transcripts"),
        ("快照 cutoff", f"{cutoff} — the newest activity in the corpus. The generator is "
                        "deterministic w.r.t. the corpus, so re-running after new work moves "
                        "the figures; both artifacts move together."),
        ("人时方法 person-hours method",
         "Union of every activity window across the transcripts; timestamps sorted, a gap "
         f"> {GAP_SECONDS // 60} minutes splits a segment; segments merged globally so "
         "overlapping sessions count once. Each segment is attributed to the phase of its start day. "
         "This is a DERIVED figure, not observed human hours."),
        ("人时合计 total", f"{grand_hours:.4f} h over {len(mined['segments'])} segments"),
        ("汇率 FX", f"CNY = USD x {USD_CNY} — the same static rate the product's own "
                    "cost::USD_CNY uses (D-20), not a live FX feed"),
        ("caveat 1", "omp only — Claude Code / OpenCode / Kimi Code sessions are NOT counted."),
        ("caveat 2", "Snapshot: the transcript corpus is frozen at the cutoff above and keeps "
                     "growing, so re-running the generator moves the numbers."),
        ("caveat 3", "P3 (DeepSeek-V4.1-Flash) reports $0.00 because its gateway publishes no "
                     "per-token price. Its token counts are real; $0.00 is a reporting gap, "
                     "not free usage."),
        ("caveat 4", "~/.omp/agent/agent.db is NOT usable (usage_cost_history is all a retired "
                     "provider with NULL timestamps; client_usage is empty). The transcript "
                     "JSONL is the authority."),
        ("caveat 5", "cacheWrite is 0 on every record — this gateway bills cache reads only."),
        ("runtime ledger", "Separate quantity: the shipped binary's own LLM spend "
                           "(Kind::Budget records, D-20). Currently empty — see "
                           "docs/expense-sheet.md §6."),
    ]:
        ws4.append(row)
    ws4.column_dimensions["A"].width = 26
    ws4.column_dimensions["B"].width = 100

    wb.save(out_path)
    _normalize_xlsx(out_path, stamp)


def _normalize_xlsx(path: str, stamp: datetime.datetime) -> None:
    """Make the emitted bytes a pure function of the mined data.

    openpyxl's writer stamps `datetime.now()` into docProps/core.xml and each zip member
    on every save, so the same data produced different bytes run to run. Rewrite both
    from the snapshot stamp — a reviewer re-running the generator gets the identical file.
    """
    import re
    import shutil
    import tempfile
    import zipfile

    iso = stamp.replace(tzinfo=None).strftime("%Y-%m-%dT%H:%M:%SZ")
    dt = stamp.timetuple()[:6]
    with zipfile.ZipFile(path) as src:
        members = [(i, src.read(i.filename)) for i in src.infolist()]
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in members:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</)",
                              rb"\g<1>" + iso.encode() + rb"\g<2>", data)
            clone = zipfile.ZipInfo(info.filename, date_time=dt)
            clone.compress_type = info.compress_type
            clone.external_attr = info.external_attr
            out.writestr(clone, data)
    shutil.move(tmp, path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sessions", default=SESSIONS)
    ap.add_argument("--json", action="store_true", help="print mined figures, write nothing")
    ap.add_argument("--xlsx-only", action="store_true", help="skip the markdown update")
    ap.add_argument("--xlsx", default=os.path.join(REPO, "docs", "expense-sheet.xlsx"))
    ap.add_argument("--md", default=os.path.join(REPO, "docs", "expense-sheet.md"))
    args = ap.parse_args()

    mined = mine(args.sessions)
    rolled = rollup(mined)
    grand_hours = sum(rolled["hours_phase"].values())
    total = totals_of(rolled)

    if args.json:
        print(json.dumps({
            "sessions": mined["sessions"], "subagents": mined["subagents"],
            "cutoff": str(mined["cutoff"]), "segments": len(mined["segments"]),
            "total_hours": round(grand_hours, 4),
            "phases": {n: dict(rolled["per_phase"].get(n, {}) | {
                "hours": round(rolled["hours_phase"].get(n, 0), 4),
                "segments": rolled["segs_phase"].get(n, 0)}) for n, _, _ in PHASES},
            "days": {d: dict(rolled["per_day"][d] | {
                "hours": round(rolled["hours_day"].get(d, 0), 4)}) for d in sorted(rolled["per_day"])},
            "models": {n: dict(mined["models"].get(n, {})) for n, _, _ in PHASES},
            "tools": {n: dict(mined["tools"].get(n, {})) for n, _, _ in PHASES},
        }, indent=1, default=str))
        return

    build_workbook(mined, rolled, args.xlsx)
    print(f"wrote {args.xlsx}")
    if not args.xlsx_only:
        update_markdown(args.md, render_markdown(mined, rolled))
        print(f"updated {args.md}")
    print(f"  {fmt_int(total['calls'])} API calls · {grand_hours:.2f} h · "
          f"${total['usd']:.4f} · cutoff {mined['cutoff']:%Y-%m-%d %H:%M} UTC")


if __name__ == "__main__":
    main()
